"""
Pruebas de licitaciones_obras.py con un feed sintético que reproduce la
estructura CODICE real de la Plataforma (rutas comprobadas sobre entradas reales).

Ejecutar desde la raíz del repositorio:
    python -m unittest discover -s tests -v
"""

import csv
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPT = os.path.join(RAIZ, "licitaciones_obras.py")
sys.path.insert(0, RAIZ)
import licitaciones_obras as L  # noqa: E402

HOY = date.today()


def dia(n):
    """Fecha relativa a hoy, para que las pruebas no caduquen."""
    return (HOY + timedelta(days=n)).isoformat()


def marca(n_dias, hora="10:00:00", tz="+02:00"):
    return f"{dia(n_dias)}T{hora}.123{tz}"


NS = ('xmlns="http://www.w3.org/2005/Atom" xmlns:at="http://purl.org/atompub/tombstones/1.0" '
      'xmlns:cac="urn:dgpe:names:draft:codice:schema:xsd:CommonAggregateComponents-2" '
      'xmlns:cbc="urn:dgpe:names:draft:codice:schema:xsd:CommonBasicComponents-2" '
      'xmlns:cac-place-ext="urn:dgpe:names:draft:codice-place-ext:schema:xsd:CommonAggregateComponents-2" '
      'xmlns:cbc-place-ext="urn:dgpe:names:draft:codice-place-ext:schema:xsd:CommonBasicComponents-2"')
BASE_ID = "https://contrataciondelestado.es/sindicacion/licitacionesPerfilContratante/"
SALIDAS = ("--salida", "licitaciones_en_plazo.csv", "--salida", "licitaciones_en_plazo.xlsx")

try:
    import openpyxl  # noqa: F401
    HAY_OPENPYXL = True
except ImportError:
    HAY_OPENPYXL = False


def entrada(n, actualizado, estado, tipo, objeto, importe, cpvs, nuts, provincia, municipio,
            limite, hora="23:59:00", terminos="", lotes="", criterios=None, duracion=("6", "MON")):
    cp = "".join(f"<cac:RequiredCommodityClassification><cbc:ItemClassificationCode>{c}"
                 f"</cbc:ItemClassificationCode></cac:RequiredCommodityClassification>" for c in cpvs)
    cr = "" if criterios is None else "<cac:AwardingTerms>" + "".join(
        f"<cac:AwardingCriteria><cbc:AwardingCriteriaTypeCode>{t}</cbc:AwardingCriteriaTypeCode>"
        f"<cbc:Description>{d}</cbc:Description><cbc:WeightNumeric>{w}</cbc:WeightNumeric></cac:AwardingCriteria>"
        for t, d, w in criterios) + "</cac:AwardingTerms>"
    loc = ((f"<cbc:CountrySubentity>{provincia}</cbc:CountrySubentity>" if provincia else "") +
           (f"<cbc:CountrySubentityCode>{nuts}</cbc:CountrySubentityCode>" if nuts else "") +
           "<cac:Address>" + (f"<cbc:CityName>{municipio}</cbc:CityName>" if municipio else "") + "</cac:Address>")
    return f"""<entry><id>{BASE_ID}{n}</id>
<link href="https://contrataciondelestado.es/wps/poc?uri=deeplink:detalle_licitacion&amp;idEvl=X{n}"/>
<title>{objeto}</title><updated>{actualizado}</updated>
<cac-place-ext:ContractFolderStatus><cbc:ContractFolderID>EXP-{n}</cbc:ContractFolderID>
<cbc-place-ext:ContractFolderStatusCode>{estado}</cbc-place-ext:ContractFolderStatusCode>
<cac-place-ext:LocatedContractingParty><cac:Party><cac:PartyName><cbc:Name>Órgano {n}</cbc:Name></cac:PartyName></cac:Party></cac-place-ext:LocatedContractingParty>
<cac:ProcurementProject><cbc:Name>{objeto}</cbc:Name><cbc:TypeCode>{tipo}</cbc:TypeCode>
<cac:BudgetAmount><cbc:EstimatedOverallContractAmount currencyID="EUR">{importe * 1.2:.2f}</cbc:EstimatedOverallContractAmount>
<cbc:TaxExclusiveAmount currencyID="EUR">{importe}</cbc:TaxExclusiveAmount></cac:BudgetAmount>{cp}
<cac:RealizedLocation>{loc}</cac:RealizedLocation>
<cac:PlannedPeriod><cbc:DurationMeasure unitCode="{duracion[1]}">{duracion[0]}</cbc:DurationMeasure></cac:PlannedPeriod>
</cac:ProcurementProject>{lotes}
<cac:TenderingTerms>{terminos}{cr}</cac:TenderingTerms>
<cac:TenderingProcess><cbc:ProcedureCode>1</cbc:ProcedureCode><cac:TenderSubmissionDeadlinePeriod>
<cbc:EndDate>{limite}</cbc:EndDate><cbc:EndTime>{hora}</cbc:EndTime></cac:TenderSubmissionDeadlinePeriod></cac:TenderingProcess>
<cac:LegalDocumentReference><cbc:ID>PCAP.pdf</cbc:ID><cac:Attachment><cac:ExternalReference>
<cbc:URI>https://contrataciondelestado.es/FileSystem/servlet/GetDocumentByIdServlet?pcap={n}</cbc:URI>
</cac:ExternalReference></cac:Attachment></cac:LegalDocumentReference>
<cac:TechnicalDocumentReference><cbc:ID>PPT.pdf</cbc:ID><cac:Attachment><cac:ExternalReference>
<cbc:URI>https://contrataciondelestado.es/FileSystem/servlet/GetDocumentByIdServlet?ppt={n}</cbc:URI>
</cac:ExternalReference></cac:Attachment></cac:TechnicalDocumentReference>
</cac-place-ext:ContractFolderStatus></entry>"""


TERMINOS_A = (
    "<cbc:PriceRevisionFormulaDescription>N/A</cbc:PriceRevisionFormulaDescription>"
    "<cac:RequiredFinancialGuarantee><cbc:GuaranteeTypeCode>3</cbc:GuaranteeTypeCode><cbc:AmountRate>5</cbc:AmountRate></cac:RequiredFinancialGuarantee>"
    "<cac:RequiredFinancialGuarantee><cbc:GuaranteeTypeCode>2</cbc:GuaranteeTypeCode><cbc:AmountRate>5</cbc:AmountRate></cac:RequiredFinancialGuarantee>"
    "<cac:TendererQualificationRequest><cac:RequiredBusinessClassificationScheme><cbc:ID>RequiredBusinessProfileCode</cbc:ID>"
    "<cac:ClassificationCategory><cbc:CodeValue>G6-2</cbc:CodeValue></cac:ClassificationCategory>"
    "<cac:ClassificationCategory><cbc:CodeValue>E1-3</cbc:CodeValue></cac:ClassificationCategory>"
    "</cac:RequiredBusinessClassificationScheme></cac:TendererQualificationRequest>")
CRITERIOS_A = [("OBJ", "Valoración económica", 60), ("OBJ", "Estabilidad laboral", 20),
               ("OBJ", "Ampliación plazo de garantía", 20)]
LOTES = "".join(
    f'<cac:ProcurementProjectLot><cbc:ID schemeName="ID_LOTE">{i}</cbc:ID><cac:TenderingTerms><cac:AwardingTerms>'
    f"<cac:AwardingCriteria><cbc:AwardingCriteriaTypeCode>OBJ</cbc:AwardingCriteriaTypeCode><cbc:Description>PRECIO</cbc:Description><cbc:WeightNumeric>70</cbc:WeightNumeric></cac:AwardingCriteria>"
    f"<cac:AwardingCriteria><cbc:AwardingCriteriaTypeCode>SUBJ</cbc:AwardingCriteriaTypeCode><cbc:Description>Memoria técnica</cbc:Description><cbc:WeightNumeric>30</cbc:WeightNumeric></cac:AwardingCriteria>"
    f"</cac:AwardingTerms></cac:TenderingTerms><cac:ProcurementProject><cbc:Name>Lote {i}</cbc:Name>"
    f"<cac:RequiredCommodityClassification><cbc:ItemClassificationCode>4523325{i}</cbc:ItemClassificationCode></cac:RequiredCommodityClassification>"
    f"</cac:ProcurementProject></cac:ProcurementProjectLot>" for i in (1, 2))


def pagina(entradas, siguiente=None, extra=""):
    enlace = f'<link rel="next" href="{siguiente}"/>' if siguiente else ""
    return f'<?xml version="1.0" encoding="UTF-8"?><feed {NS}>{enlace}{extra}{"".join(entradas)}</feed>'


# Escenario
A_ANTIGUA = entrada(1, marca(-50), "PUB", "3", "Saneamiento", 590325.73, ["45231300", "45233251"],
                    "ES618", "Sevilla", "Espartinas", dia(12), "19:00:00", TERMINOS_A, criterios=CRITERIOS_A)
A_NUEVA = A_ANTIGUA.replace(marca(-50), marca(-1)).replace(dia(12), dia(19))
B = entrada(2, marca(-2, "18:44:31"), "PUB", "3", "Pavimentación por lotes", 820000, ["45233000"],
            "ES523", "Valencia/València", "", dia(22), lotes=LOTES)
C_MIXTO = entrada(3, marca(-2, "19:09:19"), "PUB", "1", "Iluminación de Navidad", 7260, ["31500000", "45310000"],
                  "ES521", "Alicante/Alacant", "", dia(14), "09:00:00", duracion=("2", "ANN"))
D_MADRID = entrada(4, f"{dia(-2)}T12:00:00Z", "PUB", "3", "Obra en Madrid", 1500000, ["45200000"],
                   "ES300", "Madrid", "", dia(33), "14:00:00")
E_ANULADA = entrada(5, marca(-7), "PUB", "3", "Obra que se anula", 600000, ["45210000"],
                    "ES521", "Alicante/Alacant", "Elche", dia(28), "14:00:00")
G_SIN_NUTS = entrada(7, marca(-6), "PUB", "3", "Obra sin NUTS", 300000, ["45000000"], "", "Málaga", "", dia(25))
H_VENCIDA = entrada(8, marca(-6, "11:00:00"), "PUB", "3", "Obra vencida", 900000, ["45000000"],
                    "ES613", "Córdoba", "", dia(-7))
BAJA_E = (f'<at:deleted-entry ref="{BASE_ID}5" when="{marca(-2, "14:58:56")}">'
          f'<at:comment type="ANULADA"/></at:deleted-entry>')
F_NUEVA = entrada(6, marca(0, "08:00:00"), "PUB", "3", "Centro de salud", 2750000, ["45215100"],
                  "ES522", "Castellón/Castelló", "Vila-real", dia(35), "14:00:00",
                  criterios=[("OBJ", "Precio", 51), ("SUBJ", "Memoria", 49)])
B_CAMBIO = B.replace(marca(-2, "18:44:31"), marca(0, "07:00:00")).replace(">820000<", ">835000<")


class Escenario(unittest.TestCase):
    """Dos ejecuciones consecutivas: carga inicial (ZIP + feed) y actualización del día siguiente."""

    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.mkdtemp()
        os.makedirs(os.path.join(cls.dir, "zips"))
        with zipfile.ZipFile(os.path.join(cls.dir, "zips", "placsp_200001.zip"), "w") as z:
            z.writestr("licitacionesPerfilesContratanteCompleto3.atom", pagina([A_ANTIGUA]))

    def url(self, nombre):
        return Path(self.dir, nombre).resolve().as_uri()

    def escribe(self, nombre, contenido):
        with open(os.path.join(self.dir, nombre), "w", encoding="utf-8") as f:
            f.write(contenido)

    def ejecuta(self, *args):
        res = subprocess.run([sys.executable, SCRIPT, *args], cwd=self.dir, capture_output=True,
                             text=True, encoding="utf-8")
        self.assertEqual(res.returncode, 0, res.stderr)
        return res.stdout

    def csv(self, nombre="licitaciones_en_plazo.csv"):
        with open(os.path.join(self.dir, nombre), encoding="utf-8-sig") as f:
            return {r["expediente"]: r for r in csv.DictReader(f, delimiter=";")}

    def test_1_carga_inicial(self):
        self.escribe("p1.atom", pagina([A_NUEVA, B, C_MIXTO, D_MADRID], self.url("p2.atom"), BAJA_E))
        self.escribe("p2.atom", pagina([E_ANULADA, G_SIN_NUTS, H_VENCIDA]))
        self.ejecuta("--zip", "200001", "--url", self.url("p1.atom"), *SALIDAS)
        filas = self.csv()

        self.assertEqual(set(filas), {"EXP-1", "EXP-2", "EXP-7"})  # fuera: mixto, Madrid, anulada, vencida
        a = filas["EXP-1"]
        self.assertEqual(a["fecha_limite_presentacion"], f"{dia(19)} 19:00")  # versión más reciente
        self.assertEqual(a["clasificacion_exigida"], "G6-2; E1-3")
        self.assertEqual(a["garantias"], "complementaria 5%; definitiva 5%")
        self.assertEqual(a["revision_precios"], "No")
        self.assertEqual(a["pct_juicio_valor"], "0")
        self.assertEqual(a["provincia"], "Sevilla")
        self.assertEqual(a["presupuesto_base_sin_iva"], "590325,73")
        b = filas["EXP-2"]
        self.assertEqual(b["num_lotes"], "2")
        self.assertIn("[lote 1]", b["criterios_adjudicacion"])
        self.assertEqual(b["pct_juicio_valor"], "30")
        self.assertEqual(filas["EXP-7"]["comunidad"], "Andalucía")  # deducida sin NUTS
        fechas = [r["fecha_limite_presentacion"] for r in filas.values()]
        self.assertEqual(fechas, sorted(fechas))

        with open(os.path.join(self.dir, "base_licitaciones.json"), encoding="utf-8") as f:
            base = json.load(f)
        self.assertEqual(base["licitaciones"][BASE_ID + "5"]["estado"], "ANULADA")
        self.assertNotIn(BASE_ID + "4", base["licitaciones"])  # Madrid, fuera de ámbito

    def test_2_actualizacion(self):
        self.escribe("p1.atom", pagina([F_NUEVA, B_CAMBIO], self.url("p2.atom")))
        self.escribe("p2.atom", pagina([A_NUEVA, C_MIXTO, D_MADRID], self.url("p3.atom"), BAJA_E))
        self.escribe("p3.atom", pagina([E_ANULADA, G_SIN_NUTS, H_VENCIDA]))
        salida = self.ejecuta("--url", self.url("p1.atom"), *SALIDAS, "--log", "registro.log")
        filas = self.csv()

        self.assertEqual(filas["EXP-6"]["nueva"], "SÍ")
        self.assertEqual(filas["EXP-1"]["nueva"], "")
        self.assertIn("importe 820.000,00 € → 835.000,00 €", filas["EXP-2"]["cambios"])
        self.assertIn("Alcanzado el punto de la sincronización anterior", salida)

    def test_3_filtro_importe(self):
        self.ejecuta("--url", self.url("p1.atom"), "--min", "500000", "--max", "1000000",
                     "--salida", "datos/rango.csv")
        self.assertEqual(set(self.csv("datos/rango.csv")), {"EXP-1", "EXP-2"})

    @unittest.skipUnless(HAY_OPENPYXL, "openpyxl no instalado")
    def test_4_excel(self):
        """El Excel de la actualización (test_2) lleva tipos y formatos reales."""
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(self.dir, "licitaciones_en_plazo.xlsx"))
        self.assertEqual(wb.sheetnames, ["En plazo", "Resumen", "Histórico"])
        ws = wb["En plazo"]
        cab = {c.value: c.column_letter for c in ws[1]}
        self.assertIn("EnPlazo", ws.tables)
        filas = {ws[f"{cab['Expediente']}{i}"].value: i for i in range(2, ws.max_row + 1)}
        self.assertEqual(set(filas), {"EXP-1", "EXP-2", "EXP-6", "EXP-7"})
        i = filas["EXP-1"]
        importe = ws[f"{cab['Presupuesto base sin IVA']}{i}"]
        self.assertEqual(importe.value, 590325.73)
        self.assertIn("€", importe.number_format)
        limite = ws[f"{cab['Fecha límite']}{i}"]
        self.assertEqual(limite.value, datetime.strptime(f"{dia(19)} 19:00", "%Y-%m-%d %H:%M"))
        self.assertEqual(limite.number_format, "dd/mm/yyyy hh:mm")
        self.assertTrue(str(ws[f"{cab['Días restantes']}{i}"].value).startswith("=IF("))
        self.assertEqual(ws[f"{cab['Plazo de ejecución (meses)']}{i}"].value, 6)
        self.assertEqual(ws[f"{cab['Juicio de valor']}{filas['EXP-2']}"].value, 0.3)
        self.assertEqual(ws[f"{cab['Juicio de valor']}{filas['EXP-2']}"].number_format, "0%")
        self.assertEqual(ws[f"{cab['Nº lotes']}{filas['EXP-2']}"].value, 2)
        self.assertEqual(ws[f"{cab['Nueva']}{filas['EXP-6']}"].value, "SÍ")
        self.assertIsNone(ws[f"{cab['Nueva']}{i}"].value)
        pcap = ws[f"{cab['PCAP']}{i}"]
        self.assertEqual(pcap.value, "PCAP")
        self.assertTrue(pcap.hyperlink.target.startswith("https://"))
        self.assertIsInstance(ws[f"{cab['Primera detección']}{i}"].value, datetime)
        # Histórico: incluye la anulada y la vencida, con su estado
        hs = wb["Histórico"]
        hc = {c.value: c.column_letter for c in hs[1]}
        estados = {hs[f"{hc['Expediente']}{j}"].value: hs[f"{hc['Estado']}{j}"].value
                   for j in range(2, hs.max_row + 1)}
        self.assertEqual(estados["EXP-5"], "Anulada")
        self.assertEqual(estados["EXP-8"], "Plazo vencido")  # PUB en la Plataforma, pero ya venció
        self.assertEqual(estados["EXP-1"], "En plazo")
        self.assertNotIn("EXP-3", estados)  # el mixto no es contrato de obras
        # Resumen con fórmulas sobre la hoja En plazo
        rs = wb["Resumen"]
        formulas = [c.value for fila in rs.iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(any("COUNTIFS('En plazo'!" in f for f in formulas))
        # Registro de la ejecución
        with open(os.path.join(self.dir, "registro.log"), encoding="utf-8") as f:
            self.assertIn("EN PLAZO", f.read())

    def test_5_sin_openpyxl(self):
        """Sin openpyxl se genera el CSV en lugar del Excel, sin error."""
        falso = os.path.join(self.dir, "sin_openpyxl", "openpyxl")
        os.makedirs(falso, exist_ok=True)
        with open(os.path.join(falso, "__init__.py"), "w") as f:
            f.write("raise ImportError('simulado')\n")
        entorno = dict(os.environ, PYTHONPATH=os.path.dirname(falso))
        res = subprocess.run([sys.executable, SCRIPT, "--url", self.url("p1.atom"), "--salida", "solo.xlsx"],
                             cwd=self.dir, capture_output=True, text=True, encoding="utf-8", env=entorno)
        self.assertEqual(res.returncode, 0, res.stderr)
        self.assertIn("falta el paquete openpyxl", res.stdout)
        self.assertTrue(os.path.exists(os.path.join(self.dir, "solo.csv")))
        self.assertFalse(os.path.exists(os.path.join(self.dir, "solo.xlsx")))


class Utilidades(unittest.TestCase):
    def test_horario_verano_e_invierno(self):
        self.assertEqual(L.a_utc("2026-10-05T14:58:56.605+02:00"), "2026-10-05T12:58:56Z")
        self.assertEqual(L.a_utc("2026-11-05T14:58:56+01:00"), "2026-11-05T13:58:56Z")
        self.assertEqual(L.a_utc("2026-11-05T14:58:56Z"), "2026-11-05T14:58:56Z")

    def test_fecha_limite(self):
        self.assertEqual(L.limite_dt("2026-10-20 19:00"), datetime(2026, 10, 20, 19, 0))
        self.assertEqual(L.limite_dt("2026-10-20"), datetime(2026, 10, 20, 23, 59))
        self.assertIsNone(L.limite_dt(""))

    def test_plazo_meses(self):
        self.assertEqual(L.plazo_meses("6 meses"), 6)
        self.assertEqual(L.plazo_meses("2 años"), 24)
        self.assertEqual(L.plazo_meses("45 días"), 1.5)
        self.assertEqual(L.plazo_meses("12,5 meses"), 12.5)
        self.assertIsNone(L.plazo_meses(""))

    def test_ambito(self):
        reg = {"nuts": "ES521", "provincia_txt": "", "municipio": ""}
        self.assertTrue(L.en_ambito(reg, ["ES52"], []))
        self.assertFalse(L.en_ambito(reg, ["ES61"], []))
        reg = {"nuts": "", "provincia_txt": "Valencia de Alcántara", "municipio": ""}
        self.assertFalse(L.en_ambito(reg, ["ES61"], []))


if __name__ == "__main__":
    unittest.main()
