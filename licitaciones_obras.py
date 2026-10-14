#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Vigilancia de licitaciones de obras en la Plataforma de Contratación del
Sector Público (PLACSP). Versión 3.

Qué hace en cada ejecución
  1. Lee el feed ATOM de la Plataforma (y, si se pide, los ficheros ZIP mensuales).
  2. Guarda en una base local (base_licitaciones.json) todas las licitaciones de
     obras del ámbito geográfico, sin duplicados, conservando la versión más reciente
     y anotando los cambios de estado, plazo e importe.
  3. Genera licitaciones_en_plazo.xlsx (o .csv) con las que siguen abiertas, ordenadas
     por fecha límite, marcando las nuevas desde la ejecución anterior. El Excel lleva
     columnas con formato (importes, fechas, días restantes, porcentajes, enlaces),
     una hoja de resumen por provincia y el histórico completo de la base.

Uso
  python licitaciones_obras.py                      # actualización (diaria)
  python licitaciones_obras.py --meses 2            # carga inicial con los ZIP mensuales
  python licitaciones_obras.py --min 500000 --max 10000000
  python licitaciones_obras.py --nuts ES521         # solo provincia de Alicante
  python licitaciones_obras.py --salida informe.csv # CSV en lugar de Excel

Requisitos: Python 3.8+. Para el Excel, el paquete openpyxl (pip install openpyxl);
si no está instalado, se genera el CSV.
"""

import argparse
import csv
import json
import os
import re
import sys
import time
import unicodedata
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
import subprocess
import traceback
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urljoin

FEED_URL = ("https://contrataciondelestado.es/sindicacion/sindicacion_643/"
            "licitacionesPerfilesContratanteCompleto3.atom")
ZIP_URL = ("https://contrataciondelsectorpublico.gob.es/sindicacion/sindicacion_643/"
           "licitacionesPerfilesContratanteCompleto3_{}.zip")
ATOM = "{http://www.w3.org/2005/Atom}"
BASE_POR_DEFECTO = "base_licitaciones.json"
SALIDA_POR_DEFECTO = "licitaciones_en_plazo.xlsx"

# Listas de códigos de la Plataforma (si aparece un código no recogido, se exporta tal cual)
PROCEDIMIENTOS = {
    "1": "Abierto", "2": "Restringido", "3": "Negociado sin publicidad",
    "4": "Negociado con publicidad", "5": "Diálogo competitivo", "6": "Contrato menor",
    "7": "Derivado de acuerdo marco", "8": "Concurso de proyectos", "9": "Abierto simplificado",
    "10": "Asociación para la innovación", "11": "Derivado de asociación para la innovación",
    "12": "Basado en sistema dinámico de adquisición", "13": "Licitación con negociación",
    "100": "Normas internas", "999": "Otros",
}
ESTADOS = {
    "PRE": "Anuncio previo", "PUB": "En plazo", "EV": "Pendiente de adjudicación",
    "ADJ": "Adjudicada", "RES": "Resuelta", "ANUL": "Anulada", "ANULADA": "Anulada",
    "RETIRADA": "Retirada",
}
GARANTIAS = {"1": "provisional", "2": "definitiva", "3": "complementaria"}
UNIDADES = {"DAY": "días", "MON": "meses", "ANN": "años"}

COMUNIDADES = {
    "ES11": "Galicia", "ES12": "Principado de Asturias", "ES13": "Cantabria",
    "ES21": "País Vasco", "ES22": "Comunidad Foral de Navarra", "ES23": "La Rioja",
    "ES24": "Aragón", "ES30": "Comunidad de Madrid", "ES41": "Castilla y León",
    "ES42": "Castilla-La Mancha", "ES43": "Extremadura", "ES51": "Cataluña",
    "ES52": "Comunitat Valenciana", "ES53": "Illes Balears", "ES61": "Andalucía",
    "ES62": "Región de Murcia", "ES63": "Ceuta", "ES64": "Melilla", "ES70": "Canarias",
}
PROVINCIAS = {
    "ES111": "A Coruña", "ES112": "Lugo", "ES113": "Ourense", "ES114": "Pontevedra",
    "ES120": "Asturias", "ES130": "Cantabria", "ES211": "Araba/Álava", "ES212": "Gipuzkoa",
    "ES213": "Bizkaia", "ES220": "Navarra", "ES230": "La Rioja", "ES241": "Huesca",
    "ES242": "Teruel", "ES243": "Zaragoza", "ES300": "Madrid", "ES411": "Ávila",
    "ES412": "Burgos", "ES413": "León", "ES414": "Palencia", "ES415": "Salamanca",
    "ES416": "Segovia", "ES417": "Soria", "ES418": "Valladolid", "ES419": "Zamora",
    "ES421": "Albacete", "ES422": "Ciudad Real", "ES423": "Cuenca", "ES424": "Guadalajara",
    "ES425": "Toledo", "ES431": "Badajoz", "ES432": "Cáceres", "ES511": "Barcelona",
    "ES512": "Girona", "ES513": "Lleida", "ES514": "Tarragona", "ES521": "Alicante",
    "ES522": "Castellón", "ES523": "Valencia", "ES531": "Eivissa y Formentera",
    "ES532": "Mallorca", "ES533": "Menorca", "ES611": "Almería", "ES612": "Cádiz",
    "ES613": "Córdoba", "ES614": "Granada", "ES615": "Huelva", "ES616": "Jaén",
    "ES617": "Málaga", "ES618": "Sevilla", "ES620": "Murcia", "ES630": "Ceuta",
    "ES640": "Melilla", "ES703": "El Hierro", "ES704": "Fuerteventura",
    "ES705": "Gran Canaria", "ES706": "La Gomera", "ES707": "La Palma",
    "ES708": "Lanzarote", "ES709": "Tenerife",
}

CAMPOS_CSV = [
    "nueva", "expediente", "organo_contratacion", "objeto", "presupuesto_base_sin_iva",
    "valor_estimado", "fecha_limite_presentacion", "dias_restantes", "comunidad",
    "provincia", "lugar_ejecucion", "procedimiento", "plazo_ejecucion",
    "clasificacion_exigida", "criterios_adjudicacion", "pct_juicio_valor", "garantias",
    "revision_precios", "num_lotes", "cpv", "url_pcap", "url_ppt", "url_licitacion",
    "primera_deteccion", "cambios",
]


# ------------------------------------------------------------------ utilidades XML
def ln(el):
    return el.tag.rsplit("}", 1)[-1] if isinstance(el.tag, str) else ""


def ch(el, nombre):
    if el is None:
        return None
    for c in el:
        if ln(c) == nombre:
            return c
    return None


def chs(el, nombre):
    return [c for c in el if ln(c) == nombre] if el is not None else []


def ruta(el, *nombres):
    for n in nombres:
        el = ch(el, n)
        if el is None:
            return None
    return el


def tx(el, *nombres):
    e = ruta(el, *nombres) if nombres else el
    return (e.text or "").strip() if e is not None and e.text else ""


# --------------------------------------------------------------- utilidades varias
def normaliza(texto):
    texto = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in texto if not unicodedata.combining(c)).lower().strip()


def a_float(valor):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def num(v):
    """30.0 -> '30'; 12.5 -> '12,5'"""
    if v is None:
        return ""
    return str(int(v)) if float(v).is_integer() else f"{v}".replace(".", ",")


def eur(v):
    if v is None:
        return "—"
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{s} €"


_TS = re.compile(r"(\d{4})-(\d{2})-(\d{2})T(\d{2}):(\d{2}):(\d{2})(?:\.\d+)?(Z|[+-]\d{2}:\d{2})?")


def a_utc(s):
    """Marca de tiempo ATOM -> 'AAAA-MM-DDTHH:MM:SSZ' en UTC (comparable como texto)."""
    m = _TS.match(s or "")
    if not m:
        return ""
    a, mes, d, h, mi, se, tz = m.groups()
    dt = datetime(int(a), int(mes), int(d), int(h), int(mi), int(se))
    if tz and tz != "Z":
        signo = 1 if tz[0] == "+" else -1
        dt -= signo * timedelta(hours=int(tz[1:3]), minutes=int(tz[4:6]))
    return dt.strftime("%Y-%m-%dT%H:%M:%SZ")


def limite_dt(fecha):
    """'AAAA-MM-DD HH:MM' (hora local) -> datetime; sin hora = final del día."""
    if not fecha:
        return None
    try:
        if len(fecha) >= 16:
            return datetime.strptime(fecha[:16], "%Y-%m-%d %H:%M")
        return datetime.strptime(fecha[:10], "%Y-%m-%d") + timedelta(hours=23, minutes=59)
    except ValueError:
        return None


def descarga(url, destino=None, reintentos=3, timeout=120):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (licitaciones_obras.py)"})
    for intento in range(1, reintentos + 1):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                if destino is None:
                    return resp.read()
                total = int(resp.headers.get("Content-Length") or 0)
                leido = 0
                with open(destino + ".part", "wb") as f:
                    while True:
                        bloque = resp.read(1 << 20)
                        if not bloque:
                            break
                        f.write(bloque)
                        leido += len(bloque)
                        if total:
                            print(f"\r    {leido / 1e6:,.0f} / {total / 1e6:,.0f} MB", end="", flush=True)
                print()
                os.replace(destino + ".part", destino)
                return destino
        except Exception as exc:  # noqa: BLE001
            if intento == reintentos:
                raise
            print(f"  Aviso: fallo de descarga ({exc}); reintento {intento}/{reintentos - 1}...",
                  file=sys.stderr)
            time.sleep(3 * intento)


# --------------------------------------------------------------------- extracción
def extrae(entry):
    """Datos de una entrada ATOM, o None si no es una licitación."""
    cfs = ch(entry, "ContractFolderStatus")
    if cfs is None:
        return None
    pp = ch(cfs, "ProcurementProject")
    tp = ch(cfs, "TenderingProcess")
    tt = ch(cfs, "TenderingTerms")
    lotes = chs(cfs, "ProcurementProjectLot")

    cpvs = []
    for proyecto in [pp] + [ch(l, "ProcurementProject") for l in lotes]:
        for rc in chs(proyecto, "RequiredCommodityClassification"):
            c = tx(rc, "ItemClassificationCode")
            if c and c not in cpvs:
                cpvs.append(c)

    loc = ch(pp, "RealizedLocation")
    fecha = tx(tp, "TenderSubmissionDeadlinePeriod", "EndDate")
    hora = tx(tp, "TenderSubmissionDeadlinePeriod", "EndTime")
    if fecha and hora:
        fecha = f"{fecha} {hora[:5]}"

    dm = ruta(pp, "PlannedPeriod", "DurationMeasure")
    plazo = ""
    if dm is not None and (dm.text or "").strip():
        unidad = dm.get("unitCode", "")
        plazo = f"{num(a_float(dm.text.strip()))} {UNIDADES.get(unidad, unidad)}".strip()

    clasif = []
    for q in chs(tt, "TendererQualificationRequest"):
        for esquema in chs(q, "RequiredBusinessClassificationScheme"):
            for cat in chs(esquema, "ClassificationCategory"):
                v = tx(cat, "CodeValue")
                if v and v not in clasif:
                    clasif.append(v)

    criterios, peso_total, peso_jv = [], 0.0, 0.0
    lista_criterios = chs(ch(tt, "AwardingTerms"), "AwardingCriteria")
    if not lista_criterios and lotes:  # con lotes, los criterios van dentro de cada lote
        lista_criterios = chs(ruta(lotes[0], "TenderingTerms", "AwardingTerms"), "AwardingCriteria")
        if lista_criterios:
            criterios.append("[lote 1]")
    for a in lista_criterios:
        tipo = tx(a, "AwardingCriteriaTypeCode")
        peso = a_float(tx(a, "WeightNumeric"))
        etiqueta = (tx(a, "Description") or "¿?").replace("\n", " ")
        criterios.append(f"{etiqueta} {num(peso)}{' (JV)' if tipo == 'SUBJ' else ''}".strip())
        if peso:
            peso_total += peso
            if tipo == "SUBJ":
                peso_jv += peso

    garantias = []
    for g in chs(tt, "RequiredFinancialGuarantee"):
        codigo = tx(g, "GuaranteeTypeCode")
        tasa = tx(g, "AmountRate")
        if tasa:
            garantias.append(f"{GARANTIAS.get(codigo, codigo)} {num(a_float(tasa))}%")

    revision = " ".join(tx(tt, "PriceRevisionFormulaDescription").split())
    if normaliza(revision) in ("n/a", "na", "no", "-", "no procede", "no aplica"):
        revision = "No"

    link = entry.find(f"{ATOM}link")
    return {
        "id": tx(entry, "id"),
        "actualizado": a_utc(tx(entry, "updated")),
        "estado": tx(cfs, "ContractFolderStatusCode"),
        "tipo": tx(pp, "TypeCode"),
        "expediente": tx(cfs, "ContractFolderID"),
        "organo": tx(cfs, "LocatedContractingParty", "Party", "PartyName", "Name"),
        "objeto": " ".join((tx(pp, "Name") or tx(entry, "title")).split()),
        "cpv": cpvs,
        "importe": a_float(tx(pp, "BudgetAmount", "TaxExclusiveAmount")),
        "valor_estimado": a_float(tx(pp, "BudgetAmount", "EstimatedOverallContractAmount")),
        "procedimiento": tx(tp, "ProcedureCode"),
        "fecha_limite": fecha,
        "nuts": tx(loc, "CountrySubentityCode").upper(),
        "provincia_txt": tx(loc, "CountrySubentity"),
        "municipio": tx(loc, "Address", "CityName"),
        "plazo_ejecucion": plazo,
        "clasificacion": clasif,
        "criterios": criterios,
        "pct_jv": round(100 * peso_jv / peso_total, 1) if peso_total else None,
        "garantias": garantias,
        "revision_precios": revision,
        "lotes": len(lotes),
        "url_pcap": tx(cfs, "LegalDocumentReference", "Attachment", "ExternalReference", "URI"),
        "url_ppt": tx(cfs, "TechnicalDocumentReference", "Attachment", "ExternalReference", "URI"),
        "url": (link.get("href", "") if link is not None else "") or tx(entry, "id"),
    }


# -------------------------------------------------------------- ámbito geográfico
PROVINCIAS_TEXTO = {  # respaldo cuando una entrada no trae código NUTS
    "ES52": ["alicante", "alacant", "castellon", "castello", "valencia"],
    "ES61": ["almeria", "cadiz", "cordoba", "granada", "huelva", "jaen", "malaga", "sevilla"],
    "ES62": ["murcia"],
}


def en_ambito(reg, nuts_list, lugares):
    if not nuts_list and not lugares:
        return True
    cod = reg["nuts"]
    prov = normaliza(reg["provincia_txt"])
    for pref in nuts_list:
        if cod:
            if cod.startswith(pref):
                return True
        elif any(p in prov for p in PROVINCIAS_TEXTO.get(pref[:4], [])):
            return True
    texto = " | ".join(normaliza(x) for x in (reg["provincia_txt"], reg["nuts"], reg["municipio"]))
    return any(normaliza(l) in texto for l in lugares)


def es_de_obras(reg):
    """Para guardar en la base: tipo de contrato obras o algún CPV 45."""
    return reg["tipo"] == "3" or any(c.startswith("45") for c in reg["cpv"])


# ------------------------------------------------------------------- base local
class Base:
    def __init__(self, ruta_fichero):
        self.ruta = ruta_fichero
        self.datos = {"ultima_sincronizacion": "", "licitaciones": {}}
        if os.path.exists(ruta_fichero):
            with open(ruta_fichero, encoding="utf-8") as f:
                self.datos = json.load(f)
        self.nuevas, self.cambiadas = set(), set()
        self.hoy = date.today().isoformat()

    @property
    def lic(self):
        return self.datos["licitaciones"]

    def integra(self, reg):
        clave = reg["id"] or reg["expediente"]
        previo = self.lic.get(clave)
        if previo is None:
            reg["primera_deteccion"] = self.hoy
            reg["cambios"] = []
            self.lic[clave] = reg
            self.nuevas.add(clave)
            self._aplica_baja(clave)
            return
        if reg["actualizado"] <= previo.get("actualizado", ""):
            return  # ya tenemos una versión igual o más reciente
        dif = []
        if previo.get("estado") != reg["estado"]:
            dif.append(f"estado {previo.get('estado') or '—'} → {reg['estado'] or '—'}")
        if previo.get("fecha_limite") != reg["fecha_limite"]:
            dif.append(f"plazo {previo.get('fecha_limite') or '—'} → {reg['fecha_limite'] or '—'}")
        if previo.get("importe") != reg["importe"]:
            dif.append(f"importe {eur(previo.get('importe'))} → {eur(reg['importe'])}")
        if clave in self.nuevas:
            dif = []  # detectada en esta misma ejecución: no se anotan versiones intermedias
        reg["primera_deteccion"] = previo.get("primera_deteccion", self.hoy)
        reg["cambios"] = previo.get("cambios", []) + ([f"{self.hoy}: {'; '.join(dif)}"] if dif else [])
        self.lic[clave] = reg
        if dif and clave not in self.nuevas:
            self.cambiadas.add(clave)
        self._aplica_baja(clave)

    @property
    def bajas(self):
        return self.datos.setdefault("bajas", {})

    def borrada(self, ref, motivo, cuando):
        """Registra una baja (anulación/retirada). Se aplica aunque la licitación se lea después."""
        if not ref:
            return
        previa = self.bajas.get(ref)
        if previa is None or cuando > previa[0]:
            self.bajas[ref] = [cuando, motivo or "RETIRADA"]
        self._aplica_baja(ref)

    def _aplica_baja(self, ref):
        baja, r = self.bajas.get(ref), self.lic.get(ref)
        if not baja or not r:
            return
        cuando, motivo = baja
        if r.get("actualizado", "") <= cuando and r.get("estado") != motivo:
            r["cambios"] = r.get("cambios", []) + [f"{self.hoy}: {motivo.lower()} en la Plataforma"]
            r["estado"] = motivo
            if ref not in self.nuevas:
                self.cambiadas.add(ref)

    def guarda(self):
        carpeta = os.path.dirname(self.ruta)
        if carpeta:
            os.makedirs(carpeta, exist_ok=True)
        tmp = self.ruta + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.datos, f, ensure_ascii=False, indent=0)
        os.replace(tmp, self.ruta)


# --------------------------------------------------------------- procesado ATOM
class Lector:
    def __init__(self, base, nuts_list, lugares):
        self.base, self.nuts_list, self.lugares = base, nuts_list, lugares
        self.entradas = 0
        self.mas_antigua = ""
        self.mas_reciente = ""

    def procesa(self, contenido_xml):
        """Procesa una página ATOM. Devuelve (enlace_next, fecha_más_antigua_de_la_página)."""
        raiz = ET.fromstring(contenido_xml)
        antigua = ""
        for el in raiz:
            nombre = ln(el)
            if nombre == "entry":
                self.entradas += 1
                reg = extrae(el)
                if reg is None:
                    continue
                act = reg["actualizado"]
                if act:
                    antigua = act if not antigua or act < antigua else antigua
                    self.mas_antigua = act if not self.mas_antigua or act < self.mas_antigua else self.mas_antigua
                    self.mas_reciente = max(self.mas_reciente, act)
                if es_de_obras(reg) and en_ambito(reg, self.nuts_list, self.lugares):
                    self.base.integra(reg)
            elif nombre == "deleted-entry":
                comentario = ch(el, "comment")
                self.base.borrada(el.get("ref", ""),
                                  comentario.get("type", "") if comentario is not None else "",
                                  a_utc(el.get("when", "")))
        siguiente = None
        for l in raiz.findall(f"{ATOM}link"):
            if l.get("rel") == "next":
                siguiente = l.get("href")
        return siguiente, antigua


def lee_feed(lector, max_paginas, hasta, url=FEED_URL):
    """Lee el feed en vivo hasta max_paginas o hasta llegar a lo ya sincronizado."""
    paginas = 0
    while url and paginas < max_paginas:
        paginas += 1
        print(f"  Página {paginas}...", end=" ", flush=True)
        try:
            siguiente, antigua = lector.procesa(descarga(url))
        except Exception as exc:  # noqa: BLE001
            print(f"\n  ERROR al leer la página {paginas}: {exc}. Se detiene la lectura.", file=sys.stderr)
            return paginas - 1
        print(f"hasta {antigua[:16].replace('T', ' ')} UTC")
        if hasta and antigua and antigua < hasta:
            print("  Alcanzado el punto de la sincronización anterior.")
            break
        url = urljoin(url, siguiente) if siguiente else None
    return paginas


def lee_zip(lector, aaaamm, carpeta="zips"):
    os.makedirs(carpeta, exist_ok=True)
    destino = os.path.join(carpeta, f"placsp_{aaaamm}.zip")
    mes_actual = date.today().strftime("%Y%m")
    if not os.path.exists(destino) or aaaamm == mes_actual:  # el del mes en curso cambia cada día
        print(f"  Descargando fichero mensual {aaaamm}...")
        descarga(ZIP_URL.format(aaaamm), destino)
    with zipfile.ZipFile(destino) as z:
        nombres = [n for n in z.namelist() if n.lower().endswith(".atom")]
        print(f"  {aaaamm}: {len(nombres)} ficheros ATOM")
        for n in nombres:
            lector.procesa(z.read(n))


def meses_atras(n):
    d = date.today().replace(day=1)
    res = []
    for _ in range(n):
        res.append(d.strftime("%Y%m"))
        d = (d - timedelta(days=1)).replace(day=1)
    return sorted(res)


# -------------------------------------------------------------------- informe
def ubicacion(r):
    """(comunidad, provincia, lugar) a partir del código NUTS o, en su defecto, del texto."""
    nuts = r.get("nuts", "")
    comunidad = COMUNIDADES.get(nuts[:4], "")
    if not comunidad:  # sin código NUTS: se deduce del nombre de la provincia
        prov = normaliza(r.get("provincia_txt", ""))
        for cod, nombres in PROVINCIAS_TEXTO.items():
            if prov and any(n in prov for n in nombres):
                comunidad = COMUNIDADES[cod]
                break
    provincia = PROVINCIAS.get(nuts[:5], "" if len(nuts) == 4 else r.get("provincia_txt", ""))
    lugar = ", ".join(x for x in (r.get("municipio", ""), r.get("provincia_txt", "")) if x)
    lugar = f"{lugar} ({nuts})" if lugar and nuts else (lugar or nuts)
    return comunidad, provincia, lugar


def a_fila(clave, r, nuevas, ahora):
    lim = limite_dt(r["fecha_limite"])
    coma = lambda v: f"{v:.2f}".replace(".", ",") if v is not None else ""
    comunidad, provincia, lugar = ubicacion(r)
    return {
        "nueva": "SÍ" if clave in nuevas else "",
        "expediente": r["expediente"],
        "organo_contratacion": r["organo"],
        "objeto": r["objeto"],
        "presupuesto_base_sin_iva": coma(r["importe"]),
        "valor_estimado": coma(r["valor_estimado"]),
        "fecha_limite_presentacion": r["fecha_limite"],
        "dias_restantes": (lim.date() - ahora.date()).days if lim else "",
        "comunidad": comunidad,
        "provincia": provincia,
        "lugar_ejecucion": lugar,
        "procedimiento": PROCEDIMIENTOS.get(r["procedimiento"], r["procedimiento"]),
        "plazo_ejecucion": r["plazo_ejecucion"],
        "clasificacion_exigida": "; ".join(r["clasificacion"]),
        "criterios_adjudicacion": " | ".join(r["criterios"]),
        "pct_juicio_valor": num(r["pct_jv"]),
        "garantias": "; ".join(r["garantias"]),
        "revision_precios": r["revision_precios"],
        "num_lotes": r["lotes"] or "",
        "cpv": "; ".join(r["cpv"]),
        "url_pcap": r["url_pcap"],
        "url_ppt": r["url_ppt"],
        "url_licitacion": r["url"],
        "primera_deteccion": r.get("primera_deteccion", ""),
        "cambios": " | ".join(r.get("cambios", [])),
    }


def ruta_libre(ruta):
    """Nombre alternativo cuando el fichero está abierto (por ejemplo, en Excel)."""
    raiz, ext = os.path.splitext(ruta)
    return f"{raiz}_{datetime.now():%Y-%m-%d_%H%M%S}{ext}"


def escribe_csv(filas, ruta_csv):
    carpeta = os.path.dirname(ruta_csv)
    if carpeta:
        os.makedirs(carpeta, exist_ok=True)
    try:
        f = open(ruta_csv, "w", newline="", encoding="utf-8-sig")
    except PermissionError:  # el CSV está abierto en Excel
        alt = ruta_libre(ruta_csv)
        print(f"  Aviso: {ruta_csv} está abierto (¿en Excel?). Se guarda como {alt}")
        ruta_csv, f = alt, open(alt, "w", newline="", encoding="utf-8-sig")
    with f:
        w = csv.DictWriter(f, fieldnames=CAMPOS_CSV, delimiter=";")
        w.writeheader()
        w.writerows(filas)
    return ruta_csv


# ------------------------------------------------------------------ salida Excel
def plazo_meses(texto):
    """'6 meses' -> 6; '2 años' -> 24; '45 días' -> 1.5 (días / 30)."""
    m = re.match(r"\s*([\d.,]+)\s*(\S+)", texto or "")
    if not m:
        return None
    n = a_float(m.group(1).replace(",", "."))
    unidad = normaliza(m.group(2))
    if n is None:
        return None
    if unidad.startswith("mes"):
        return n
    if unidad.startswith("ano"):
        return n * 12
    if unidad.startswith("dia"):
        return round(n / 30, 1)
    return None


def utc_a_local(marca):
    try:
        dt = datetime.strptime(marca, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
        return dt.astimezone().replace(tzinfo=None)
    except (TypeError, ValueError):
        return None


def fecha_iso(texto):
    try:
        return datetime.strptime(texto[:10], "%Y-%m-%d")
    except (TypeError, ValueError):
        return None


# Columnas: (cabecera, ancho, tipo, valor). Tipos: txt, eur, fechahora, fecha, ent, dec, pct,
# dias (fórmula sobre la columna "Fecha límite"), y enlace:<texto visible>.
def _columnas_en_plazo(nuevas):
    return [
        ("Nueva", 7, "txt", lambda k, r: "SÍ" if k in nuevas else None),
        ("Fecha límite", 16, "fechahora", lambda k, r: limite_dt(r.get("fecha_limite", ""))),
        ("Días restantes", 9, "dias", None),
        ("Presupuesto base sin IVA", 16, "eur", lambda k, r: r.get("importe")),
        ("Valor estimado", 16, "eur", lambda k, r: r.get("valor_estimado")),
        ("Comunidad", 20, "txt", lambda k, r: ubicacion(r)[0]),
        ("Provincia", 12, "txt", lambda k, r: ubicacion(r)[1]),
        ("Municipio", 20, "txt", lambda k, r: r.get("municipio")),
        ("Órgano de contratación", 40, "txt", lambda k, r: r.get("organo")),
        ("Objeto", 70, "txt", lambda k, r: r.get("objeto")),
        ("Expediente", 20, "txt", lambda k, r: r.get("expediente")),
        ("Procedimiento", 18, "txt", lambda k, r: PROCEDIMIENTOS.get(r.get("procedimiento", ""), r.get("procedimiento"))),
        ("Plazo de ejecución (meses)", 11, "dec", lambda k, r: plazo_meses(r.get("plazo_ejecucion", ""))),
        ("Clasificación exigida", 16, "txt", lambda k, r: "; ".join(r.get("clasificacion", []))),
        ("Juicio de valor", 9, "pct", lambda k, r: None if r.get("pct_jv") is None else r["pct_jv"] / 100),
        ("Criterios de adjudicación", 60, "txt", lambda k, r: " | ".join(r.get("criterios", []))),
        ("Garantías", 22, "txt", lambda k, r: "; ".join(r.get("garantias", []))),
        ("Revisión de precios", 18, "txt", lambda k, r: r.get("revision_precios")),
        ("Nº lotes", 7, "ent", lambda k, r: r.get("lotes") or None),
        ("CPV", 28, "txt", lambda k, r: "; ".join(r.get("cpv", []))),
        ("PCAP", 7, "enlace:PCAP", lambda k, r: r.get("url_pcap")),
        ("PPT", 7, "enlace:PPT", lambda k, r: r.get("url_ppt")),
        ("Ficha PLACSP", 9, "enlace:Ver", lambda k, r: r.get("url")),
        ("Primera detección", 12, "fecha", lambda k, r: fecha_iso(r.get("primera_deteccion", ""))),
        ("Cambios", 50, "txt", lambda k, r: " | ".join(r.get("cambios", []))),
    ]


def estado_texto(r, ahora):
    """Estado legible. Una licitación PUB con el plazo ya vencido figura como 'Plazo vencido'
    hasta que la Plataforma actualice su estado."""
    lim = limite_dt(r.get("fecha_limite", ""))
    if r.get("estado") == "PUB" and lim and lim < ahora:
        return "Plazo vencido"
    return ESTADOS.get(r.get("estado", ""), r.get("estado"))


def _columnas_historico(ahora):
    return [
        ("Estado", 22, "txt", lambda k, r: estado_texto(r, ahora)),
        ("Fecha límite", 16, "fechahora", lambda k, r: limite_dt(r.get("fecha_limite", ""))),
        ("Presupuesto base sin IVA", 16, "eur", lambda k, r: r.get("importe")),
        ("Valor estimado", 16, "eur", lambda k, r: r.get("valor_estimado")),
        ("Comunidad", 20, "txt", lambda k, r: ubicacion(r)[0]),
        ("Provincia", 12, "txt", lambda k, r: ubicacion(r)[1]),
        ("Municipio", 20, "txt", lambda k, r: r.get("municipio")),
        ("Órgano de contratación", 40, "txt", lambda k, r: r.get("organo")),
        ("Objeto", 70, "txt", lambda k, r: r.get("objeto")),
        ("Expediente", 20, "txt", lambda k, r: r.get("expediente")),
        ("Procedimiento", 18, "txt", lambda k, r: PROCEDIMIENTOS.get(r.get("procedimiento", ""), r.get("procedimiento"))),
        ("Plazo de ejecución (meses)", 11, "dec", lambda k, r: plazo_meses(r.get("plazo_ejecucion", ""))),
        ("Clasificación exigida", 16, "txt", lambda k, r: "; ".join(r.get("clasificacion", []))),
        ("Juicio de valor", 9, "pct", lambda k, r: None if r.get("pct_jv") is None else r["pct_jv"] / 100),
        ("Nº lotes", 7, "ent", lambda k, r: r.get("lotes") or None),
        ("CPV", 28, "txt", lambda k, r: "; ".join(r.get("cpv", []))),
        ("Ficha PLACSP", 9, "enlace:Ver", lambda k, r: r.get("url")),
        ("PCAP", 7, "enlace:PCAP", lambda k, r: r.get("url_pcap")),
        ("Primera detección", 12, "fecha", lambda k, r: fecha_iso(r.get("primera_deteccion", ""))),
        ("Última actualización en PLACSP", 16, "fechahora", lambda k, r: utc_a_local(r.get("actualizado", ""))),
        ("Cambios", 50, "txt", lambda k, r: " | ".join(r.get("cambios", []))),
    ]


FORMATOS = {"eur": '#,##0.00 "€"', "fechahora": "dd/mm/yyyy hh:mm", "fecha": "dd/mm/yyyy",
            "ent": "0", "dec": "0.0", "pct": "0%", "dias": "0"}


def escribe_xlsx(ruta, abiertas, historico, nuevas, meta):
    """Libro con tres hojas: En plazo, Resumen e Histórico. Requiere openpyxl."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    letra = Font(name="Arial", size=10)
    negrita = Font(name="Arial", size=10, bold=True)
    enlace = Font(name="Arial", size=10, color="0563C1", underline="single")
    cabecera = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    relleno_cab = PatternFill("solid", fgColor="1F4E78")
    rojo = (PatternFill("solid", fgColor="FFC7CE"), Font(name="Arial", size=10, color="9C0006", bold=True))
    ambar = (PatternFill("solid", fgColor="FFEB9C"), Font(name="Arial", size=10, color="9C5700"))
    verde = (PatternFill("solid", fgColor="C6EFCE"), Font(name="Arial", size=10, color="006100", bold=True))
    naranja = PatternFill("solid", fgColor="FCE4D6")

    def limpia(v):
        return ILLEGAL_CHARACTERS_RE.sub("", v)[:32000] if isinstance(v, str) else v

    def hoja_tabla(ws, nombre_tabla, columnas, registros):
        """Escribe una tabla de Excel con formato por columna. Devuelve {cabecera: letra}."""
        letras = {c[0]: get_column_letter(i + 1) for i, c in enumerate(columnas)}
        col_fecha = letras.get("Fecha límite")
        for j, (titulo, ancho, _, _) in enumerate(columnas, 1):
            c = ws.cell(row=1, column=j, value=titulo)
            c.font, c.fill = cabecera, relleno_cab
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = ancho
        ws.row_dimensions[1].height = 32
        for i, (clave, r) in enumerate(registros, 2):
            for j, (titulo, _, tipo, valor) in enumerate(columnas, 1):
                c = ws.cell(row=i, column=j)
                c.font = letra
                if tipo == "dias":
                    c.value = f'=IF({col_fecha}{i}="","",INT({col_fecha}{i})-TODAY())'
                elif tipo.startswith("enlace:"):
                    url = valor(clave, r)
                    if url:
                        c.value, c.hyperlink, c.font = tipo.split(":", 1)[1], url, enlace
                else:
                    v = limpia(valor(clave, r))
                    c.value = v if v not in ("", None) else None
                    if isinstance(v, str) and v.startswith("="):
                        c.data_type = "s"  # texto, no fórmula
                if tipo in FORMATOS:
                    c.number_format = FORMATOS[tipo]
        n = len(registros)
        ultima = get_column_letter(len(columnas))
        if n:
            tabla = Table(displayName=nombre_tabla, ref=f"A1:{ultima}{n + 1}")
            tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tabla)
        else:
            ws.cell(row=2, column=1, value="Sin licitaciones que cumplan los filtros.").font = letra
        return letras, n

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True  # días restantes y resumen al día al abrir

    # --- Hoja 1: En plazo
    ws = wb.active
    ws.title = "En plazo"
    letras, n = hoja_tabla(ws, "EnPlazo", _columnas_en_plazo(nuevas), abiertas)
    ws.freeze_panes = "D2"
    if n:
        ult = n + 1
        d = letras["Días restantes"]
        ws.conditional_formatting.add(f"{d}2:{d}{ult}", FormulaRule(
            formula=[f'AND(ISNUMBER({d}2),{d}2<=7)'], fill=rojo[0], font=rojo[1], stopIfTrue=True))
        ws.conditional_formatting.add(f"{d}2:{d}{ult}", FormulaRule(
            formula=[f'AND(ISNUMBER({d}2),{d}2<=14)'], fill=ambar[0], font=ambar[1], stopIfTrue=True))
        a = letras["Nueva"]
        ws.conditional_formatting.add(f"{a}2:{a}{ult}", FormulaRule(
            formula=[f'{a}2="SÍ"'], fill=verde[0], font=verde[1]))
        cb = letras["Cambios"]
        ws.conditional_formatting.add(f"{cb}2:{cb}{ult}", FormulaRule(
            formula=[f'LEN({cb}2)>0'], fill=naranja))

    # --- Hoja 2: Resumen (fórmulas sobre la hoja En plazo)
    rs = wb.create_sheet("Resumen")
    rs.column_dimensions["A"].width = 34
    rs.column_dimensions["B"].width = 16
    for col in "CDEF":
        rs.column_dimensions[col].width = 16
    t = rs.cell(row=1, column=1, value="Licitaciones de obras en plazo · Resumen")
    t.font = Font(name="Arial", size=14, bold=True, color="1F4E78")
    fin = max(n, 1) + 1
    rng = lambda cab: f"'En plazo'!${letras[cab]}$2:${letras[cab]}${fin}"
    filas_meta = [
        ("Actualizado", meta["actualizado"], "dd/mm/yyyy hh:mm"),
        ("Ámbito (códigos NUTS)", meta["ambito"], None),
        ("Rango de presupuesto base sin IVA", meta["rango"], None),
        ("Fuentes leídas", meta["fuentes"], None),
        ("Entradas leídas", meta["entradas"], "#,##0"),
        ("Periodo cubierto en esta lectura", meta["periodo"], None),
        ("Licitaciones en la base", meta["en_base"], "#,##0"),
        ("En plazo", f"=COUNTA({rng('Objeto')})", "#,##0"),
        ("Importe total en plazo", f"=SUM({rng('Presupuesto base sin IVA')})", FORMATOS["eur"]),
        ("Nuevas desde la ejecución anterior", f'=COUNTIF({rng("Nueva")},"SÍ")', "#,##0"),
        ("Con cambios registrados", f'=COUNTIF({rng("Cambios")},"?*")', "#,##0"),
        ("Vencen en 7 días o menos", f'=COUNTIFS({rng("Días restantes")},">=0",{rng("Días restantes")},"<=7")', "#,##0"),
    ]
    for i, (etq, v, fmt) in enumerate(filas_meta, 3):
        rs.cell(row=i, column=1, value=etq).font = negrita
        c = rs.cell(row=i, column=2, value=v)
        c.font = letra
        c.alignment = Alignment(horizontal="left")
        if fmt:
            c.number_format = fmt

    fila = 3 + len(filas_meta) + 1
    rs.cell(row=fila, column=1, value="Por comunidad y provincia").font = Font(name="Arial", size=11, bold=True, color="1F4E78")
    fila += 1
    for j, cab in enumerate(["Comunidad", "Provincia", "Nº en plazo", "Importe", "Nuevas", "Vencen ≤ 7 días"], 1):
        c = rs.cell(row=fila, column=j, value=cab)
        c.font, c.fill = cabecera, relleno_cab
    primera = fila + 1
    grupos = sorted({(ubicacion(r)[0], ubicacion(r)[1]) for _, r in abiertas})
    for comunidad, provincia in grupos:
        fila += 1
        crit_c = f'"{comunidad}"'
        crit_p = f'"{provincia}"'
        cond = f'{rng("Comunidad")},{crit_c},{rng("Provincia")},{crit_p}'
        valores = [
            comunidad or "(sin comunidad)",
            provincia or "(sin provincia)",
            f"=COUNTIFS({cond})",
            f"=SUMIFS({rng('Presupuesto base sin IVA')},{cond})",
            f'=COUNTIFS({cond},{rng("Nueva")},"SÍ")',
            f'=COUNTIFS({cond},{rng("Días restantes")},">=0",{rng("Días restantes")},"<=7")',
        ]
        for j, v in enumerate(valores, 1):
            c = rs.cell(row=fila, column=j, value=v)
            c.font = letra
            c.number_format = FORMATOS["eur"] if j == 4 else ("#,##0" if j > 2 else "General")
    if grupos:
        fila += 1
        rs.cell(row=fila, column=1, value="Total").font = negrita
        for j, col in ((3, "C"), (4, "D"), (5, "E"), (6, "F")):
            c = rs.cell(row=fila, column=j, value=f"=SUM({col}{primera}:{col}{fila - 1})")
            c.font = negrita
            c.number_format = FORMATOS["eur"] if j == 4 else "#,##0"
    fila += 2
    notas = [
        "Fuente: Plataforma de Contratación del Sector Público (feed ATOM licitacionesPerfilesContratanteCompleto3) "
        "y ficheros mensuales de datos abiertos del Ministerio de Hacienda.",
        "Los días restantes y los vencimientos se recalculan con la fecha del día al abrir el libro.",
        "Plazo de ejecución en meses: los plazos en días se dividen entre 30 y los de años se multiplican por 12.",
        "Este libro se regenera en cada ejecución: las anotaciones que se hagan en él se pierden. "
        "Para el seguimiento, copiar las filas al libro de preselección.",
    ]
    for nota in notas:
        rs.cell(row=fila, column=1, value=nota).font = Font(name="Arial", size=9, italic=True, color="595959")
        fila += 1

    # --- Hoja 3: Histórico
    hs = wb.create_sheet("Histórico")
    _, nh = hoja_tabla(hs, "Historico", _columnas_historico(meta["actualizado"]), historico)
    hs.freeze_panes = "C2"

    for hoja in wb:  # impresión: horizontal, ajustada al ancho y con la cabecera en cada página
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 0
    for hoja in (ws, hs):
        hoja.print_title_rows = "1:1"

    carpeta = os.path.dirname(ruta)
    if carpeta:
        os.makedirs(carpeta, exist_ok=True)
    try:
        wb.save(ruta)
    except PermissionError:  # el libro está abierto en Excel
        alt = ruta_libre(ruta)
        print(f"  Aviso: {ruta} está abierto (¿en Excel?). Se guarda como {alt}")
        wb.save(alt)
        ruta = alt
    return ruta


# ------------------------------------------------------------------- registro
def abre_registro(ruta):
    if os.path.exists(ruta) and os.path.getsize(ruta) > 2_000_000:
        os.replace(ruta, ruta + ".1")
    f = open(ruta, "a", encoding="utf-8")
    f.write(f"\n===== {datetime.now():%Y-%m-%d %H:%M:%S} =====\n")
    return f


class Registro:
    """Duplica la salida en un fichero de registro (útil en la ejecución programada)."""

    def __init__(self, fichero, consola):
        self.f, self.consola = fichero, consola

    def write(self, texto):
        self.f.write(texto)
        self.f.flush()
        if self.consola is not None:
            try:
                self.consola.write(texto)
            except Exception:  # noqa: BLE001
                pass
        return len(texto)

    def flush(self):
        self.f.flush()
        if self.consola is not None:
            try:
                self.consola.flush()
            except Exception:  # noqa: BLE001
                pass


def abre(ruta):
    try:
        if hasattr(os, "startfile"):
            os.startfile(os.path.abspath(ruta))  # Windows
        elif sys.platform == "darwin":
            subprocess.Popen(["open", ruta])
        else:
            subprocess.Popen(["xdg-open", ruta])
    except Exception as exc:  # noqa: BLE001
        print(f"  Aviso: no se pudo abrir {ruta}: {exc}")


# ----------------------------------------------------------------------- main
def main():
    for flujo in (sys.stdout, sys.stderr):  # evita errores de codificación en Windows al redirigir
        try:
            flujo.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--nuts", default="ES52,ES61",
                    help="Prefijos NUTS separados por comas (por defecto ES52,ES61: C. Valenciana y Andalucía)")
    ap.add_argument("--lugar", default="", help="Además, texto de provincia o municipio, separado por comas")
    ap.add_argument("--paginas", type=int, default=30,
                    help="Máximo de páginas del feed en vivo (se detiene antes si llega a lo ya leído)")
    ap.add_argument("--meses", type=int, default=0,
                    help="Carga los N últimos ficheros ZIP mensuales (carga inicial; recomendado 2)")
    ap.add_argument("--zip", action="append", default=[], metavar="AAAAMM",
                    help="Carga un fichero mensual concreto; se puede repetir")
    ap.add_argument("--min", type=float, default=None, help="Presupuesto mínimo sin IVA")
    ap.add_argument("--max", type=float, default=None, help="Presupuesto máximo sin IVA")
    ap.add_argument("--incluir-mixtos", action="store_true",
                    help="Incluir contratos que no son de tipo obras pero tienen algún CPV 45")
    ap.add_argument("--base", default=BASE_POR_DEFECTO, help="Fichero de la base local")
    ap.add_argument("--salida", action="append", default=None, metavar="FICHERO",
                    help=f"Fichero de resultados .xlsx o .csv; se puede repetir (por defecto {SALIDA_POR_DEFECTO})")
    ap.add_argument("--log", default="", metavar="FICHERO",
                    help="Añade la salida a un fichero de registro (ejecución programada)")
    ap.add_argument("--abrir", choices=["nunca", "siempre", "novedades"], default="nunca",
                    help="Abrir el resultado al terminar: nunca, siempre o solo si hay nuevas o cambios")
    ap.add_argument("--url", default=FEED_URL, help=argparse.SUPPRESS)
    args = ap.parse_args()
    if args.log:  # con pythonw no hay consola: sys.stdout es None y solo se escribe el registro
        fichero = abre_registro(args.log)
        sys.stdout, sys.stderr = Registro(fichero, sys.stdout), Registro(fichero, sys.stderr)
    try:
        ejecuta(args)
    except Exception:  # noqa: BLE001
        print("ERROR no previsto:\n" + traceback.format_exc())
        sys.exit(1)


def ejecuta(args):
    salidas = args.salida or [SALIDA_POR_DEFECTO]
    nuts_list = [x.strip().upper() for x in args.nuts.split(",") if x.strip()]
    lugares = [x.strip() for x in args.lugar.split(",") if x.strip()]
    base = Base(args.base)
    primera_vez = not base.lic
    lector = Lector(base, nuts_list, lugares)
    inicio = datetime.now()

    zips = sorted(set(args.zip + meses_atras(args.meses))) if (args.zip or args.meses) else []
    for z in zips:
        try:
            lee_zip(lector, z)
        except Exception as exc:  # noqa: BLE001
            print(f"  ERROR con el fichero mensual {z}: {exc}", file=sys.stderr)

    print("Leyendo el feed en vivo...")
    # margen de 1 día sobre la última sincronización para no perder nada
    hasta = ""
    if base.datos.get("ultima_sincronizacion"):
        ult = datetime.strptime(base.datos["ultima_sincronizacion"], "%Y-%m-%dT%H:%M:%SZ")
        hasta = (ult - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M:%SZ")
    paginas = lee_feed(lector, args.paginas, hasta, args.url)

    if lector.mas_reciente:
        base.datos["ultima_sincronizacion"] = max(base.datos.get("ultima_sincronizacion", ""),
                                                  lector.mas_reciente)
    base.guarda()

    # Licitaciones abiertas
    ahora = datetime.now()
    abiertas = []
    for clave, r in base.lic.items():
        if r.get("estado") != "PUB":
            continue
        lim = limite_dt(r["fecha_limite"])
        if lim and lim < ahora:
            continue
        if not args.incluir_mixtos and r.get("tipo") != "3":
            continue
        if args.min is not None and (r["importe"] is None or r["importe"] < args.min):
            continue
        if args.max is not None and (r["importe"] is None or r["importe"] > args.max):
            continue
        abiertas.append((clave, r))
    abiertas.sort(key=lambda kr: kr[1]["fecha_limite"] or "9999")

    # Resumen
    nuevas_abiertas = [r for k, r in abiertas if k in base.nuevas]
    cambiadas_abiertas = [r for k, r in abiertas if k in base.cambiadas]
    total = sum(r["importe"] or 0 for _, r in abiertas)
    rango = ""
    if args.min is not None or args.max is not None:
        rango = f" entre {eur(args.min) if args.min is not None else '0 €'} y {eur(args.max) if args.max is not None else '∞'}"

    # Resultados
    escritos = []
    for salida in salidas:
        if salida.lower().endswith(".xlsx"):
            try:
                import openpyxl  # noqa: F401
            except ImportError:
                alt = salida[:-5] + ".csv"
                print(f"  Aviso: falta el paquete openpyxl (pip install openpyxl). Se genera {alt} en su lugar.")
                if alt not in salidas:
                    escritos.append(escribe_csv([a_fila(k, r, base.nuevas, ahora) for k, r in abiertas], alt))
                continue
            historico = [(k, r) for k, r in base.lic.items() if args.incluir_mixtos or r.get("tipo") == "3"]
            historico.sort(key=lambda kr: kr[1].get("actualizado", ""), reverse=True)
            periodo = ""
            if lector.mas_antigua:
                a, b = utc_a_local(lector.mas_antigua), utc_a_local(lector.mas_reciente)
                periodo = f"{a:%d/%m/%Y} – {b:%d/%m/%Y}" if a and b else ""
            meta = {
                "actualizado": ahora.replace(microsecond=0),
                "ambito": ", ".join(nuts_list + lugares) or "Toda España",
                "rango": rango.strip() or "Sin límite",
                "fuentes": f"{len(zips)} fichero(s) mensual(es) + {paginas} página(s) del feed",
                "entradas": lector.entradas,
                "periodo": periodo or "Sin entradas nuevas",
                "en_base": len(base.lic),
            }
            escritos.append(escribe_xlsx(salida, abiertas, historico, base.nuevas, meta))
        else:
            escritos.append(escribe_csv([a_fila(k, r, base.nuevas, ahora) for k, r in abiertas], salida))

    print("\n" + "=" * 78)
    print(f"Fuentes:  {len(zips)} fichero(s) mensual(es) + {paginas} página(s) del feed "
          f"({(datetime.now() - inicio).seconds} s)")
    print(f"Entradas leídas en total:     {lector.entradas}")
    if lector.mas_antigua:
        print(f"Periodo cubierto:             {lector.mas_antigua[:10]} a {lector.mas_reciente[:10]}")
    print(f"Licitaciones en la base:      {len(base.lic)}")
    if primera_vez:
        print("Primera ejecución: todas las licitaciones cuentan como nuevas.")
    print(f"\nEN PLAZO (obras{rango}):  {len(abiertas)}  |  Importe total: {eur(total)}")
    print(f"  Nuevas desde la última ejecución:  {len(nuevas_abiertas)}")
    print(f"  Con cambios (plazo, importe...):   {len(cambiadas_abiertas)}")
    for ruta in escritos:
        print(f"Resultado guardado en: {ruta}")

    urgentes = [r for _, r in abiertas if limite_dt(r["fecha_limite"]) and
                (limite_dt(r["fecha_limite"]).date() - ahora.date()).days <= 7]
    if urgentes:
        print("\nVencen en 7 días o menos:")
        for r in urgentes:
            print(f"  {r['fecha_limite']}  {eur(r['importe']):>16}  {r['organo'][:40]} | {r['objeto'][:50]}")
    if not primera_vez and nuevas_abiertas:
        print("\nNuevas (por importe):")
        for r in sorted(nuevas_abiertas, key=lambda r: r["importe"] or 0, reverse=True)[:10]:
            print(f"  {eur(r['importe']):>16}  plazo {r['fecha_limite']}  {r['organo'][:40]} | {r['objeto'][:50]}")
    if cambiadas_abiertas:
        print("\nCambios detectados:")
        for r in cambiadas_abiertas:
            print(f"  {r['expediente']}: {r['cambios'][-1]}")
    print("=" * 78)

    hay_novedades = bool(nuevas_abiertas or cambiadas_abiertas)
    if escritos and (args.abrir == "siempre" or (args.abrir == "novedades" and hay_novedades)):
        abre(escritos[0])


if __name__ == "__main__":
    main()
